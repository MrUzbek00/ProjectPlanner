#!/usr/bin/env python3
"""Build the Word and Excel deliverable package from project Markdown.

The deliverable documents are authored as Markdown under
``projects/<project>/deliverables/``. This script converts each file to a
formatted ``.docx`` and collects the tables tagged with an ``xlsx`` marker into
workbooks.

A table is routed to a workbook by placing an HTML comment directly above it::

    <!-- xlsx: workbook=requirements; sheet=Functional Requirements -->

    | FR ID | Requirement | ... |
    | --- | --- | ... |

Usage::

    python tools/build_deliverables.py projects/<project>/deliverables
    python tools/build_deliverables.py projects/<project>/deliverables --check
    python tools/build_deliverables.py <dir> --out <dir> --project "Name"

``--check`` validates without writing anything and reports unresolved content.
"""

from __future__ import annotations

import argparse
import datetime as _dt
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

# --------------------------------------------------------------------------
# Configuration
# --------------------------------------------------------------------------

#: Workbook key -> output file name suffix. Unknown keys fall back to the key.
WORKBOOK_FILES = {
    "requirements": "Requirements-and-Traceability",
    "plan": "Project-Plan-and-Roadmap",
}

#: Preferred document order when several Markdown files are present.
DOCUMENT_ORDER = [
    "brd",
    "prd",
    "srs",
    "sow",
    "user-journey",
    "wireframe",
    "project-plan",
    "package",
]

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
NUMBER_RE = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
HR_RE = re.compile(r"^\s*(\*\s*\*\s*\*|-\s*-\s*-|_\s*_\s*_)[\s*\-_]*$")
XLSX_MARKER_RE = re.compile(r"<!--\s*xlsx:(?P<body>.*?)-->", re.S)
IMAGE_RE = re.compile(r"^!\[(?P<caption>.*?)\]\((?P<src>[^)]+)\)\s*$")
COMMENT_OPEN_RE = re.compile(r"^\s*<!--")
COMMENT_CLOSE_RE = re.compile(r"-->\s*$")
TBD_RE = re.compile(r"TBD(\s*\(Q-\d+\))?")

INLINE_RE = re.compile(
    r"(?P<code>`[^`]+`)"
    r"|(?P<bold>\*\*[^*]+\*\*)"
    r"|(?P<italic>(?<!\*)\*(?!\*)[^*]+\*(?!\*)|(?<![\w_])_[^_]+_(?![\w_]))"
    r"|(?P<link>\[[^\]]+\]\([^)]*\))"
)


# --------------------------------------------------------------------------
# Document model
# --------------------------------------------------------------------------


@dataclass
class Block:
    kind: str  # heading | paragraph | table | code | list | quote | rule | image
    text: str = ""
    src: str = ""
    level: int = 0
    rows: list[list[str]] = field(default_factory=list)
    items: list[tuple[int, str]] = field(default_factory=list)
    ordered: bool = False
    language: str = ""
    sheet: dict | None = None
    heading_path: tuple[str, ...] = ()


@dataclass
class Document:
    path: Path
    meta: dict[str, str]
    blocks: list[Block]

    @property
    def title(self) -> str:
        if self.meta.get("title"):
            return self.meta["title"]
        for block in self.blocks:
            if block.kind == "heading" and block.level == 1:
                return block.text
        return self.path.stem.replace("-", " ").title()


# --------------------------------------------------------------------------
# Parsing
# --------------------------------------------------------------------------


def parse_meta(text: str) -> tuple[dict[str, str], str]:
    """Extract the ``doc-meta`` comment block and return it with the remainder."""
    match = re.search(r"<!--\s*doc-meta\s*(?P<body>.*?)-->", text, re.S)
    if not match:
        return {}, text
    meta: dict[str, str] = {}
    for line in match.group("body").splitlines():
        line = line.strip()
        if not line or ":" not in line:
            continue
        key, _, value = line.partition(":")
        meta[key.strip().lower()] = value.strip()
    return meta, text[: match.start()] + text[match.end():]


def parse_sheet_marker(raw: str) -> dict:
    """Parse ``workbook=x; sheet=Y; freeze=B2`` into a dictionary."""
    spec: dict[str, str] = {}
    for part in raw.split(";"):
        part = part.strip()
        if not part or "=" not in part:
            continue
        key, _, value = part.partition("=")
        spec[key.strip().lower()] = value.strip()
    return spec


def split_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    cells: list[str] = []
    current = ""
    escaped = False
    for char in line:
        if escaped:
            current += char
            escaped = False
        elif char == "\\":
            escaped = True
        elif char == "|":
            cells.append(current.strip())
            current = ""
        else:
            current += char
    cells.append(current.strip())
    return cells


def parse_markdown(path: Path) -> Document:
    raw = path.read_text(encoding="utf-8")
    meta, body = parse_meta(raw)
    lines = body.splitlines()
    blocks: list[Block] = []
    pending_sheet: dict | None = None
    heading_path: list[str] = []
    index = 0

    def current_path() -> tuple[str, ...]:
        return tuple(heading_path)

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        # Fenced code block.
        if stripped.startswith("```"):
            language = stripped[3:].strip()
            index += 1
            buffer: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                buffer.append(lines[index])
                index += 1
            index += 1
            blocks.append(
                Block(
                    kind="code",
                    text="\n".join(buffer),
                    language=language,
                    heading_path=current_path(),
                )
            )
            continue

        # HTML comment, possibly spanning several lines.
        if COMMENT_OPEN_RE.match(line):
            buffer = [line]
            while index < len(lines) and not COMMENT_CLOSE_RE.search(lines[index]):
                index += 1
                if index < len(lines):
                    buffer.append(lines[index])
            index += 1
            comment = "\n".join(buffer)
            marker = XLSX_MARKER_RE.search(comment)
            if marker:
                pending_sheet = parse_sheet_marker(marker.group("body"))
            continue

        heading = HEADING_RE.match(line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            del heading_path[level - 1:]
            while len(heading_path) < level - 1:
                heading_path.append("")
            heading_path.append(text)
            blocks.append(
                Block(kind="heading", text=text, level=level, heading_path=current_path())
            )
            index += 1
            continue

        # Table: a pipe row followed by a separator row.
        if (
            stripped.startswith("|")
            and index + 1 < len(lines)
            and TABLE_SEP_RE.match(lines[index + 1].strip())
        ):
            header = split_table_row(lines[index])
            index += 2
            rows = [header]
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            blocks.append(
                Block(
                    kind="table",
                    rows=rows,
                    sheet=pending_sheet,
                    heading_path=current_path(),
                )
            )
            pending_sheet = None
            continue

        image = IMAGE_RE.match(stripped)
        if image:
            blocks.append(
                Block(
                    kind="image",
                    text=image.group("caption").strip(),
                    src=image.group("src").strip(),
                    heading_path=current_path(),
                )
            )
            index += 1
            continue

        if HR_RE.match(line):
            blocks.append(Block(kind="rule", heading_path=current_path()))
            index += 1
            continue

        if stripped.startswith(">"):
            buffer = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                buffer.append(lines[index].strip().lstrip(">").strip())
                index += 1
            blocks.append(
                Block(kind="quote", text=" ".join(buffer), heading_path=current_path())
            )
            continue

        if BULLET_RE.match(line) or NUMBER_RE.match(line):
            items: list[tuple[int, str]] = []
            ordered = bool(NUMBER_RE.match(line))
            while index < len(lines):
                candidate = lines[index]
                bullet = BULLET_RE.match(candidate)
                number = NUMBER_RE.match(candidate)
                if bullet:
                    indent = len(bullet.group(1)) // 2
                    items.append((indent, bullet.group(2).strip()))
                elif number:
                    indent = len(number.group(1)) // 2
                    items.append((indent, number.group(3).strip()))
                elif candidate.strip() and candidate.startswith(("   ", "\t")) and items:
                    depth, text = items[-1]
                    items[-1] = (depth, f"{text} {candidate.strip()}")
                else:
                    break
                index += 1
            blocks.append(
                Block(
                    kind="list",
                    items=items,
                    ordered=ordered,
                    heading_path=current_path(),
                )
            )
            continue

        # Paragraph.
        buffer = []
        while index < len(lines) and lines[index].strip():
            candidate = lines[index]
            if (
                HEADING_RE.match(candidate)
                or candidate.strip().startswith(("|", ">", "```"))
                or COMMENT_OPEN_RE.match(candidate)
                or BULLET_RE.match(candidate)
                or NUMBER_RE.match(candidate)
                or IMAGE_RE.match(candidate.strip())
            ):
                break
            buffer.append(candidate.strip())
            index += 1
        if buffer:
            blocks.append(
                Block(
                    kind="paragraph",
                    text=" ".join(buffer),
                    heading_path=current_path(),
                )
            )
        else:
            index += 1

    return Document(path=path, meta=meta, blocks=blocks)


def plain_text(value: str) -> str:
    """Strip inline Markdown syntax for spreadsheet cells."""
    value = re.sub(r"\[([^\]]+)\]\(([^)]*)\)", r"\1", value)
    value = value.replace("**", "").replace("`", "")
    value = re.sub(r"(?<![\w*])\*(?!\*)([^*]+)\*(?!\*)", r"\1", value)
    value = value.replace("<br>", "\n").replace("\\|", "|")
    return value.strip()


# --------------------------------------------------------------------------
# Word output
# --------------------------------------------------------------------------


def _require(module: str, package: str):
    try:
        return __import__(module)
    except ImportError:  # pragma: no cover - environment dependent
        raise SystemExit(
            f"Missing dependency '{package}'. Install the document tooling with:\n"
            f"    python -m pip install -r tools/requirements-docs.txt"
        )


def _add_field(paragraph, instruction: str) -> None:
    """Insert a Word field code (used for the table of contents and page numbers)."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, end):
        run._r.append(element)


def _shade(cell_or_paragraph, color: str) -> None:
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), color)
    cell_element = getattr(cell_or_paragraph, "_tc", None)
    if cell_element is not None:
        properties = cell_element.get_or_add_tcPr()
    else:
        properties = cell_or_paragraph._p.get_or_add_pPr()
    properties.append(shading)


def _write_runs(paragraph, text: str, mono: bool = False) -> None:
    """Render inline Markdown into Word runs."""
    from docx.shared import Pt

    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        if match.group("code"):
            run = paragraph.add_run(match.group("code")[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif match.group("bold"):
            paragraph.add_run(match.group("bold")[2:-2]).bold = True
        elif match.group("italic"):
            paragraph.add_run(match.group("italic")[1:-1]).italic = True
        elif match.group("link"):
            label = re.match(r"\[([^\]]+)\]", match.group("link")).group(1)
            run = paragraph.add_run(label)
            run.underline = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])
    if mono:
        for run in paragraph.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(9)


def build_docx(document: Document, out_path: Path, project: str) -> None:
    docx = _require("docx", "python-docx")
    from docx import Document as WordDocument
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Inches, RGBColor

    word = WordDocument()

    normal = word.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    section = word.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    for margin in ("left_margin", "right_margin"):
        setattr(section, margin, Inches(0.8))
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)

    meta = document.meta
    title = meta.get("title") or document.title
    project_name = meta.get("project") or project

    # Cover page.
    spacer = word.add_paragraph()
    spacer.paragraph_format.space_after = Pt(48)

    heading = word.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0x1F, 0x36, 0x4D)

    if meta.get("subtitle"):
        subtitle = word.add_paragraph()
        run = subtitle.add_run(meta["subtitle"])
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x5A, 0x63, 0x6E)

    cover_rows = [
        ("Project", project_name),
        ("Client / department", meta.get("client", "")),
        ("Version", meta.get("version", "")),
        ("Date", meta.get("date", "")),
        ("Author", meta.get("author", "")),
        ("Status", meta.get("status", "")),
        ("Language", meta.get("language", "")),
        ("Source records", meta.get("source", "")),
    ]
    cover_rows = [(label, value) for label, value in cover_rows if value]
    if cover_rows:
        word.add_paragraph()
        table = word.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for label, value in cover_rows:
            cells = table.add_row().cells
            cells[0].paragraphs[0].add_run(label).bold = True
            cells[1].text = value
        for row in table.rows:
            row.cells[0].width = Inches(1.9)
            row.cells[1].width = Inches(4.3)

    word.add_page_break()

    # Table of contents.
    toc_heading = word.add_paragraph()
    run = toc_heading.add_run("Contents")
    run.bold = True
    run.font.size = Pt(16)
    toc = word.add_paragraph()
    _add_field(toc, r'TOC \o "1-3" \h \z \u')
    note = word.add_paragraph()
    run = note.add_run(
        "Open in Word and refresh fields (Ctrl+A, then F9) to populate this table."
    )
    run.italic = True
    run.font.size = Pt(8.5)
    word.add_page_break()

    # Page numbers in the footer.
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"{title} — ").font.size = Pt(8)
    _add_field(footer, "PAGE")
    footer.add_run(" / ").font.size = Pt(8)
    _add_field(footer, "NUMPAGES")
    for run in footer.runs:
        run.font.size = Pt(8)

    content_width = section.page_width - section.left_margin - section.right_margin

    skip_first_h1 = True
    for block in document.blocks:
        if block.kind == "heading":
            if block.level == 1 and skip_first_h1:
                skip_first_h1 = False
                continue
            level = min(block.level, 4)
            paragraph = word.add_heading("", level=level)
            _write_runs(paragraph, block.text)
            continue

        if block.kind == "paragraph":
            paragraph = word.add_paragraph()
            _write_runs(paragraph, block.text)
            continue

        if block.kind == "quote":
            paragraph = word.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.3)
            _write_runs(paragraph, block.text)
            for run in paragraph.runs:
                run.italic = True
            continue

        if block.kind == "rule":
            word.add_paragraph()
            continue

        if block.kind == "list":
            for depth, text in block.items:
                style = "List Number" if block.ordered else "List Bullet"
                paragraph = word.add_paragraph(style=style)
                paragraph.paragraph_format.left_indent = Inches(0.25 + 0.25 * depth)
                paragraph.paragraph_format.space_after = Pt(2)
                _write_runs(paragraph, text)
            continue

        if block.kind == "code":
            paragraph = word.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.left_indent = Inches(0.1)
            for line_index, line in enumerate(block.text.splitlines()):
                run = paragraph.add_run(line)
                run.font.name = "Consolas"
                run.font.size = Pt(8)
                if line_index < len(block.text.splitlines()) - 1:
                    run.add_break()
            _shade(paragraph, "F4F5F7")
            continue

        if block.kind == "image":
            _write_image(word, block, document.path.parent, content_width)
            continue

        if block.kind == "table" and block.rows:
            _write_table(word, block)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    word.save(str(out_path))


def _write_image(word, block: Block, base: Path, content_width) -> None:
    """Embed a diagram export and render its caption beneath it."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    target = (base / block.src).resolve()
    holder = word.add_paragraph()
    holder.alignment = WD_ALIGN_PARAGRAPH.CENTER
    holder.paragraph_format.space_after = Pt(2)

    if target.is_file():
        run = holder.add_run()
        run.add_picture(str(target))
        shape = word.inline_shapes[-1]
        # Shrink a wide diagram to the text column; never upscale a small one.
        if shape.width > content_width:
            ratio = content_width / shape.width
            shape.width = content_width
            shape.height = int(shape.height * ratio)
    else:
        run = holder.add_run(f"[missing image: {block.src}]")
        run.bold = True
        run.font.color.rgb = RGBColor(0xB8, 0x54, 0x50)

    if block.text:
        caption = word.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(10)
        _write_runs(caption, block.text)
        for run in caption.runs:
            run.italic = True
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0x5A, 0x63, 0x6E)


def _write_table(word, block: Block) -> None:
    from docx.shared import Pt, Inches

    header, *body = block.rows
    columns = max(len(row) for row in block.rows)
    table = word.add_table(rows=1, cols=columns)
    table.style = "Table Grid"
    table.autofit = True

    if columns > 10:
        font_size = Pt(6.5)
    elif columns > 6:
        font_size = Pt(7.5)
    else:
        font_size = Pt(9)

    for index in range(columns):
        cell = table.rows[0].cells[index]
        text = header[index] if index < len(header) else ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(plain_text(text))
        run.bold = True
        run.font.size = font_size
        _shade(cell, "E8EDF3")

    for row_values in body:
        cells = table.add_row().cells
        for index in range(columns):
            value = row_values[index] if index < len(row_values) else ""
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            _write_runs(paragraph, value)
            for run in paragraph.runs:
                run.font.size = font_size

    word.add_paragraph().paragraph_format.space_after = Pt(4)


# --------------------------------------------------------------------------
# Excel output
# --------------------------------------------------------------------------


INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def safe_sheet_name(name: str, used: Iterable[str]) -> str:
    cleaned = INVALID_SHEET_CHARS.sub("-", name).strip() or "Sheet"
    cleaned = cleaned[:31]
    used = set(used)
    if cleaned not in used:
        return cleaned
    counter = 2
    while True:
        suffix = f" ({counter})"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        if candidate not in used:
            return candidate
        counter += 1


NUMERIC_RE = re.compile(r"^-?\d{1,3}(,\d{3})*(\.\d+)?$|^-?\d+(\.\d+)?$")


def cell_value(raw: str):
    """Return a number for plain numeric cells so Excel can sort and total them."""
    text = plain_text(raw)
    if not text:
        return None
    if NUMERIC_RE.match(text):
        cleaned = text.replace(",", "")
        try:
            return int(cleaned) if "." not in cleaned else float(cleaned)
        except ValueError:
            return text
    # Values carrying a unit, a currency symbol or a percent sign stay as text so
    # the exported figure reads exactly as the source record wrote it.
    return text


def collect_sheets(documents: list[Document]) -> dict[str, list[tuple[Document, Block]]]:
    workbooks: dict[str, list[tuple[Document, Block]]] = {}
    for document in documents:
        for block in document.blocks:
            if block.kind == "table" and block.sheet:
                key = block.sheet.get("workbook", "requirements")
                workbooks.setdefault(key, []).append((document, block))
    return workbooks


def build_xlsx(key: str, entries: list[tuple[Document, Block]], out_path: Path, project: str) -> list[str]:
    _require("openpyxl", "openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)

    header_fill = PatternFill("solid", fgColor="1F364D")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    body_font = Font(size=10)
    index_fill = PatternFill("solid", fgColor="E8EDF3")
    thin = Side(style="thin", color="BFC7D1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    index_sheet = workbook.create_sheet("Index")
    index_sheet.append(["Sheet", "Source document", "Source section", "Columns", "Data rows"])
    created: list[str] = []

    for document, block in entries:
        name = block.sheet.get("sheet") or "Sheet"
        sheet_name = safe_sheet_name(name, workbook.sheetnames)
        sheet = workbook.create_sheet(sheet_name)
        created.append(sheet_name)

        header, *body = block.rows
        columns = max(len(row) for row in block.rows)
        header = [plain_text(header[i]) if i < len(header) else "" for i in range(columns)]
        sheet.append(header)

        for row_values in body:
            sheet.append(
                [cell_value(row_values[i]) if i < len(row_values) else None for i in range(columns)]
            )

        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border

        for index in range(1, columns + 1):
            letter = get_column_letter(index)
            longest = max(
                (len(str(cell.value)) for cell in sheet[letter] if cell.value is not None),
                default=10,
            )
            sheet.column_dimensions[letter].width = max(12, min(48, longest + 2))

        sheet.row_dimensions[1].height = 30
        freeze = block.sheet.get("freeze", "A2")
        sheet.freeze_panes = freeze
        if sheet.max_row >= 1:
            sheet.auto_filter.ref = f"A1:{get_column_letter(columns)}{max(sheet.max_row, 1)}"

        section = " > ".join(part for part in block.heading_path if part) or "-"
        index_sheet.append(
            [sheet_name, document.path.name, section, columns, max(sheet.max_row - 1, 0)]
        )

    index_sheet["A1"].value = "Sheet"
    for cell in index_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
    for row in index_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate([28, 30, 46, 10, 12], start=1):
        index_sheet.column_dimensions[get_column_letter(index)].width = width
    index_sheet.freeze_panes = "A2"
    index_sheet.sheet_properties.tabColor = "1F364D"

    title_row = index_sheet.max_row + 2
    index_sheet.cell(row=title_row, column=1, value="Project").font = Font(bold=True, size=10)
    index_sheet.cell(row=title_row, column=2, value=project).font = body_font
    index_sheet.cell(row=title_row + 1, column=1, value="Workbook").font = Font(bold=True, size=10)
    index_sheet.cell(row=title_row + 1, column=2, value=key).font = body_font
    index_sheet.cell(row=title_row + 2, column=1, value="Generated").font = Font(bold=True, size=10)
    index_sheet.cell(
        row=title_row + 2, column=2, value=_dt.date.today().isoformat()
    ).font = body_font
    index_sheet.cell(row=title_row, column=1).fill = index_fill

    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(out_path))
    return created


# --------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------


def validate(documents: list[Document]) -> tuple[list[str], list[str]]:
    warnings: list[str] = []
    errors: list[str] = []

    for document in documents:
        name = document.path.name
        tables = [b for b in document.blocks if b.kind == "table"]
        empty_tables = [b for b in tables if len(b.rows) <= 1]
        tbd_count = sum(
            len(TBD_RE.findall(b.text)) for b in document.blocks if b.kind == "paragraph"
        ) + sum(
            len(TBD_RE.findall(" ".join(cell for row in b.rows for cell in row)))
            for b in tables
        )

        if not document.meta:
            warnings.append(f"{name}: no doc-meta block; the cover page will be sparse.")
        if empty_tables:
            warnings.append(
                f"{name}: {len(empty_tables)} of {len(tables)} tables have no data rows."
            )
        if tbd_count:
            warnings.append(f"{name}: {tbd_count} unresolved TBD entries.")
        if "TEMPLATE" in " ".join(
            b.text for b in document.blocks[:6] if b.kind == "paragraph"
        ):
            warnings.append(
                f"{name}: still carries template instruction text; remove it before release."
            )

        images = [b for b in document.blocks if b.kind == "image"]
        for block in images:
            if not (document.path.parent / block.src).is_file():
                errors.append(f"{name}: referenced image not found: {block.src}")
            elif not block.text.strip():
                warnings.append(f"{name}: image {block.src} has no caption.")

        for block in document.blocks:
            if block.kind == "table" and block.sheet:
                if not block.sheet.get("sheet"):
                    errors.append(f"{name}: xlsx marker without a sheet name.")
                widths = {len(row) for row in block.rows}
                if len(widths) > 1:
                    warnings.append(
                        f"{name}: table '{block.sheet.get('sheet')}' has uneven column counts {sorted(widths)}."
                    )

    return warnings, errors


# --------------------------------------------------------------------------
# Entry point
# --------------------------------------------------------------------------


def document_sort_key(path: Path) -> tuple[int, str]:
    stem = path.stem.lower()
    for index, token in enumerate(DOCUMENT_ORDER):
        if token in stem:
            return index, stem
    return len(DOCUMENT_ORDER), stem


def slugify(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", value).strip("-")
    return slug or "Project"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("source", type=Path, help="Directory holding the deliverable Markdown files")
    parser.add_argument("--out", type=Path, default=None, help="Output directory (default: <source>/build)")
    parser.add_argument("--project", default=None, help="Project name used on cover pages and file names")
    parser.add_argument("--check", action="store_true", help="Validate only; write nothing")
    parser.add_argument("--only", default=None, help="Build a single Markdown file by name")
    args = parser.parse_args(argv)

    source: Path = args.source
    if not source.is_dir():
        print(f"error: {source} is not a directory", file=sys.stderr)
        return 2

    paths = sorted(
        (p for p in source.glob("*.md") if not p.name.startswith("_")),
        key=document_sort_key,
    )
    if args.only:
        paths = [p for p in paths if args.only.lower() in p.name.lower()]
    if not paths:
        print(f"error: no Markdown files found in {source}", file=sys.stderr)
        return 2

    documents = [parse_markdown(path) for path in paths]
    project = args.project or next(
        (d.meta.get("project") for d in documents if d.meta.get("project") and d.meta["project"] != "TBD"),
        source.resolve().parent.name,
    )

    warnings, errors = validate(documents)
    for message in warnings:
        print(f"warning: {message}")
    for message in errors:
        print(f"error: {message}", file=sys.stderr)

    if args.check:
        print(
            f"\nchecked {len(documents)} document(s): "
            f"{len(warnings)} warning(s), {len(errors)} error(s)"
        )
        return 1 if errors else 0

    if errors:
        print("\nrefusing to build while errors remain", file=sys.stderr)
        return 1

    out_dir: Path = args.out or (source / "build")
    out_dir.mkdir(parents=True, exist_ok=True)
    prefix = slugify(project)

    for document in documents:
        target = out_dir / f"{prefix}-{document.path.stem}.docx"
        build_docx(document, target, project)
        print(f"wrote {target}")

    workbooks = collect_sheets(documents)
    if not workbooks:
        print("note: no tables carried an xlsx marker; no workbook was produced.")
    for key, entries in sorted(workbooks.items()):
        suffix = WORKBOOK_FILES.get(key, slugify(key))
        target = out_dir / f"{prefix}-{suffix}.xlsx"
        sheets = build_xlsx(key, entries, target, project)
        print(f"wrote {target} ({len(sheets)} sheet(s): {', '.join(sheets)})")

    print(
        f"\nbuilt {len(documents)} document(s) and {len(workbooks)} workbook(s) into {out_dir}"
    )
    print("open each exported file and verify it before delivery")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
